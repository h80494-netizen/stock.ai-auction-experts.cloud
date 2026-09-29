import urllib.request
import io
import re
import pandas as pd
import yfinance as yf
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_kospi_market_breadth_excel(start_date_target="2013-08-13", end_date_target="2026-09-22", output_file="KOSPI_Market_Breadth_2013_2026.xlsx"):
    print("=== 1. KOSPI 전 종목 리스트 수집 (KRX KIND) ===")
    url = "http://kind.krx.co.kr/corpgeneral/corpList.do?method=download&searchType=13&marketType=stockMkt"
    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
    
    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req) as resp:
        html = resp.read().decode('euc-kr', errors='ignore')
        df_kind = pd.read_html(io.StringIO(html))[0]
        
    yf_tickers = []
    for code in df_kind['종목코드']:
        code_str = str(code).strip()
        match = re.search(r'\d+', code_str)
        if match:
            clean_code = match.group(0).zfill(6)
            yf_tickers.append(f"{clean_code}.KS")
            
    yf_tickers = list(set(yf_tickers))
    print(f"  -> KOSPI 대상 종목 수: {len(yf_tickers)}개\n")

    print(f"=== 2. KOSPI 일별 종가 시세 다운로드 (2012-08-01 ~ {end_date_target}) ===")
    start_date = "2012-08-01"
    
    data = yf.download(yf_tickers, start=start_date, end="2026-09-25", progress=False, auto_adjust=True)
    
    if data.empty or 'Close' not in data:
        print("[오류] 시세 데이터를 가져오지 못했습니다.")
        return

    price_matrix = data['Close'].dropna(how='all')
    print(f"  -> 수집된 거래일 수: {len(price_matrix)}일\n")

    print("=== 3. 일자별 등락 종목 수(상승/하락/보합) 및 ADR / AD Line 산출 중 ===")
    price_diff = price_matrix.diff()
    
    up_cnt = (price_diff > 0).sum(axis=1)
    down_cnt = (price_diff < 0).sum(axis=1)
    same_cnt = (price_diff == 0).sum(axis=1)
    
    net_diff = up_cnt - down_cnt
    ad_line = net_diff.cumsum()
    
    rolling_up = up_cnt.rolling(window=20).sum()
    rolling_down = down_cnt.rolling(window=20).sum()
    adr_20 = (rolling_up / rolling_down) * 100

    df_breadth = pd.DataFrame({
        '일자': pd.to_datetime(price_matrix.index).strftime('%Y-%m-%d'),
        '상승종목수': up_cnt.values,
        '하락종목수': down_cnt.values,
        '등락차(상승-하락)': net_diff.values,
        '누적등락선(AD Line)': ad_line.values,
        'ADR_20일(%)': adr_20.values,
        '보합종목수': same_cnt.values
    })

    # 2013-08-13 ~ 2026-09-22 기간 슬라이싱 및 정렬
    df_result = df_breadth[(df_breadth['일자'] >= start_date_target) & (df_breadth['일자'] <= end_date_target)].copy()
    df_result = df_result.sort_values('일자').reset_index(drop=True)

    cols_order = [
        "일자", "상승종목수", "하락종목수", "등락차(상승-하락)", 
        "누적등락선(AD Line)", "ADR_20일(%)", "보합종목수"
    ]
    df_result = df_result[cols_order]

    # openpyxl 서식 적용 엑셀 저장
    print(f"\n=== 4. 엑셀 고급 서식 적용 및 파일 저장 ({end_date_target} 까지) ===")
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df_result.to_excel(writer, sheet_name="KOSPI_등락통계", index=False)
        ws = writer.sheets["KOSPI_등락통계"]

        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for col_idx in range(1, len(cols_order) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for row in range(2, len(df_result) + 2):
            for col in range(1, len(cols_order) + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = Font(name="맑은 고딕", size=10)
                cell.border = thin_border
                
                col_name = cols_order[col - 1]
                if col_name == "일자":
                    cell.alignment = center_align
                elif col_name == "ADR_20일(%)":
                    cell.alignment = right_align
                    cell.number_format = '#,##0.0'
                else:
                    cell.alignment = right_align
                    cell.number_format = '#,##0'

        ws.freeze_panes = "B2"

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    csv_file = output_file.replace(".xlsx", ".csv")
    df_result.to_csv(csv_file, index=False)

    print(f"\n★ [Market Breadth] 수집 완결 ({df_result['일자'].min()} ~ {df_result['일자'].max()})")
    print(f"★ 파일 저장 완료: [{output_file}], [{csv_file}] (총 {len(df_result)}개 거래일)")

if __name__ == "__main__":
    generate_kospi_market_breadth_excel()
