import os
import sys
import datetime
import requests
import pandas as pd
import numpy as np
import openpyxl
from bs4 import BeautifulSoup
import yfinance as yf

EXCEL_PATH = r"c:\Users\llll\Documents\두인경매\주식투자\data\선물가격업데이트.xlsx"
HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'}

def get_last_valid_date_and_row(excel_path):
    """
    Returns last valid date row index where authentic historical data is present.
    Row 652 (2013-08-12) is the last authentic historical entry before generated rows.
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active
    
    last_row = 652
    last_date = datetime.date(2013, 8, 12)
    wb.close()
    return last_row, last_date

def fetch_yfinance_ohlcv(symbol, start_date, end_date):
    """
    Fetches genuine OHLCV data from yfinance.
    Returns DataFrame indexed by date or empty DataFrame if unavailable.
    """
    try:
        s_str = start_date.strftime("%Y-%m-%d")
        e_str = (end_date + datetime.timedelta(days=1)).strftime("%Y-%m-%d")
        df = yf.Ticker(symbol).history(start=s_str, end=e_str)
        if not df.empty:
            df.index = pd.to_datetime(df.index).date
        return df
    except Exception as e:
        print(f"[Info] yfinance fetch error for {symbol}: {e}")
        return pd.DataFrame()

def fetch_naver_k200_index(pages=15):
    """
    Fetches authentic KOSPI 200 Index daily records from Naver Finance.
    """
    records = []
    for page in range(1, pages + 1):
        url = f"https://finance.naver.com/sise/sise_index_day.naver?code=KPI200&page={page}"
        try:
            res = requests.get(url, headers=HEADERS, timeout=5)
            soup = BeautifulSoup(res.text, 'html.parser')
            rows = soup.select('table.type_1 tr')
            for tr in rows:
                cols = tr.find_all('td')
                if len(cols) >= 6:
                    date_text = cols[0].text.strip()
                    if date_text and '.' in date_text:
                        dt = datetime.datetime.strptime(date_text, "%Y.%m.%d").date()
                        close_val = float(cols[1].text.replace(',', ''))
                        records.append({'Date': dt, 'Close': close_val})
        except Exception:
            pass
    df = pd.DataFrame(records)
    if not df.empty:
        df = df.drop_duplicates(subset=['Date']).set_index('Date').sort_index()
    return df

def safe_val(val, default=None):
    if pd.isna(val) or val is None:
        return default
    if isinstance(val, (pd.Series, np.ndarray)):
        val = val[0] if len(val) > 0 else default
        if pd.isna(val) or val is None:
            return default
    return val

def update_excel_data():
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"Target Excel file not found: {EXCEL_PATH}")

    wb = openpyxl.load_workbook(EXCEL_PATH)
    sheet = wb.active

    max_r = sheet.max_row
    if max_r > 652:
        print(f"Cleaning up unverified rows from Row 653 to {max_r}...")
        for r in range(653, max_r + 1):
            for c in range(1, 41):
                sheet.cell(row=r, column=c, value=None)

    last_row = 652
    last_date = datetime.date(2013, 8, 12)

    start_date = last_date + datetime.timedelta(days=1)
    end_date = datetime.date.today()

    print(f"Updating data strictly without artificial estimates: {start_date} ~ {end_date}")

    df_usdkrw = fetch_yfinance_ohlcv("USDKRW=X", start_date, end_date)
    df_kodex = fetch_yfinance_ohlcv("122630.KS", start_date, end_date)
    df_k200_naver = fetch_naver_k200_index(pages=15)

    dates = []
    curr = start_date
    while curr <= end_date:
        if curr.weekday() < 5:
            dates.append(curr)
        curr += datetime.timedelta(days=1)

    print(f"Processing {len(dates)} business days...")

    current_row = last_row + 1

    for dt in dates:
        dt_str = dt.strftime("%Y-%m-%d")

        # 1. KOSPI 200 (Cols B~F)
        if not df_k200_naver.empty and dt in df_k200_naver.index:
            k_close = safe_val(df_k200_naver.loc[dt]['Close'])
        else:
            k_close = None

        k_open = k_high = k_low = k_vol = None

        # 2. USD/KRW (Cols G~K)
        if not df_usdkrw.empty and dt in df_usdkrw.index:
            r_fx = df_usdkrw.loc[dt]
            fx_open = safe_val(r_fx['Open'])
            fx_high = safe_val(r_fx['High'])
            fx_low = safe_val(r_fx['Low'])
            fx_close = safe_val(r_fx['Close'])
            fx_vol = safe_val(r_fx['Volume'])
            if fx_vol is not None:
                fx_vol = int(fx_vol)
        else:
            fx_open = fx_high = fx_low = fx_close = fx_vol = None

        # 3. Adjusted KOSPI 200 (Cols L~O)
        if k_close is not None and fx_close is not None and fx_close > 0:
            adj_ratio = fx_close / 1000.0
            adj_close = round(k_close / adj_ratio, 2)
        else:
            adj_close = None
        adj_open = adj_high = adj_low = None

        # 4. KODEX Leverage (Col AM)
        if not df_kodex.empty and dt in df_kodex.index:
            k_val = safe_val(df_kodex.loc[dt]['Close'])
            kodex_close = int(k_val) if k_val is not None else None
        else:
            kodex_close = None

        # 40 Columns row (Strict Authentic / Blank for missing)
        row_vals = [
            dt_str,             # A: Date
            k_open, k_high, k_low, k_close, k_vol, # B~F: KOSPI200
            fx_open, fx_high, fx_low, fx_close, fx_vol, # G~K: USDKRW
            adj_open, adj_high, adj_low, adj_close, # L~O: Adj KOSPI200
            None,               # P: Foreign Net Buy
            None,               # Q: Call Theo
            None,               # R: Put Theo
            None,               # S: KTB Low
            None,               # T: KTB Close
            None,               # U: Kos200 Inst Net Buy
            None,               # V: Open Interest (Futures)
            None,               # W: Program Net Buy
            None,               # X: Arbitrage Net Buy
            None,               # Y: Up count
            None,               # Z: Down count
            None, None, None, None, # AA~AD: Call ATM (O, H, L, C)
            None, None, None, None, # AE~AH: Put ATM (O, H, L, C)
            None, None, None, None, # AI~AL: Call/Put Net Buys
            kodex_close,        # AM: KODEX Leverage
            None                # AN: Futures Close
        ]

        for col_idx, val in enumerate(row_vals, 1):
            sheet.cell(row=current_row, column=col_idx, value=val)

        current_row += 1

    wb.save(EXCEL_PATH)
    wb.close()
    print(f"Successfully cleaned up estimated data and updated authentic records up to row {current_row - 1}")

if __name__ == "__main__":
    update_excel_data()
