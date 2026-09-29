import os
import openpyxl
import pandas as pd

data_dir = r"C:\Users\llll\Documents\두인경매\주식투자\data"
files_to_check = [
    os.path.join(data_dir, "선물가격업데이트_260917.xlsx"),
    os.path.join(data_dir, "선물가격업데이트.xlsx"),
    r"C:\Users\llll\Documents\두인경매\주식투자\backend\테트리스U7_240808.xlsm",
    os.path.join(data_dir, "기술지표10000_쵝오_모니터링_20130801_외인수정.xlsx"),
]

for filepath in files_to_check:
    if os.path.exists(filepath):
        print(f"\n==========================================")
        print(f"Checking: {os.path.basename(filepath)}")
        try:
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            print("Sheets:", wb.sheetnames)
            sheet = wb.active
            rows = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i < 5:
                    rows.append(row)
                else:
                    break
            print("First 3 rows:")
            for r in rows[:3]:
                print(r[:25])
        except Exception as e:
            print(f"Error reading with openpyxl: {e}")
            try:
                df = pd.read_excel(filepath, nrows=3)
                print("Pandas Columns:", df.columns.tolist()[:30])
            except Exception as pe:
                print(f"Error reading with pandas: {pe}")
